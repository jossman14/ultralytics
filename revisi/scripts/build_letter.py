"""Render response_to_reviewers.md as .docx, resolving the same {{tokens}} the manuscript uses.

The letter quotes measured numbers back to the reviewers, so it reads from
tokens_filled.json rather than carrying its own copies that could drift from the
manuscript after a re-run.
"""
import json
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt

ROOT = Path("/home/ftib/ultralytics")
SRC = ROOT / "revisi/response_to_reviewers.md"
DST = ROOT / "revisi/out/response_to_reviewers.docx"


def runs(par, text):
    for chunk in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**"):
            par.add_run(chunk[2:-2]).bold = True
        elif chunk.startswith("*") and chunk.endswith("*"):
            par.add_run(chunk[1:-1]).italic = True
        else:
            par.add_run(chunk)


def main():
    tok = json.loads((ROOT / "revisi/out/tokens_filled.json").read_text())
    src = SRC.read_text()
    for k, v in tok.items():
        src = src.replace("{{" + k + "}}", v)
    left = re.findall(r"\{\{[^}]+\}\}", src)
    if left:
        raise SystemExit(f"unresolved tokens in letter: {sorted(set(left))}")

    d = Document()
    d.styles["Normal"].font.name = "Times New Roman"
    d.styles["Normal"].font.size = Pt(11)

    for block in re.split(r"\n\s*\n", src):
        b = block.strip()
        if not b:
            continue
        if b.startswith("|"):
            rows = [r for r in b.splitlines() if r.strip().startswith("|")]
            cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
            cells = [c for c in cells if not all(set(x) <= set("-: ") for x in c)]
            tbl = d.add_table(rows=len(cells), cols=len(cells[0]))
            tbl.style = "Table Grid"
            for i, row in enumerate(cells):
                for j, c in enumerate(row):
                    runs(tbl.cell(i, j).paragraphs[0], c)
            continue
        m = re.match(r"^(#+)\s*(.*)", b)
        if m:
            d.add_heading(m.group(2), level=min(len(m.group(1)), 4))
            continue
        runs(d.add_paragraph(), " ".join(b.split()))

    d.save(DST)
    print(f"wrote {DST}")


if __name__ == "__main__":
    main()
